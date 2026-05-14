from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
OUT_DIR = DOCS_DIR / "word"

DOC_SOURCES = [
    ("README.md", "RestaurantOS Documentation Index", "Documentation Pack Index"),
    ("UAT_PLAN.md", "RestaurantOS UAT Plan", "User Acceptance Testing"),
    ("USER_MANUAL.md", "RestaurantOS User Manual", "Operational User Guide"),
    ("FLOW_DIAGRAMS.md", "RestaurantOS Flow Diagrams", "Business and System Process Flows"),
    ("SEQUENCE_DIAGRAMS.md", "RestaurantOS Sequence Diagrams", "Interaction and Sync Sequences"),
    ("SYSTEM_ARCHITECTURE.md", "RestaurantOS System Architecture", "Technical Architecture Document"),
    ("ON_PREMISE_DEPLOYMENT.md", "RestaurantOS On-Premise Deployment Guide", "Docker Installation, Operations, Backup, and Restore"),
]

TITLE_COLOR = RGBColor(31, 78, 121)
ACCENT_COLOR = "1F4E79"
LIGHT_FILL = "EAF2F8"
TABLE_HEADER_FILL = "D9EAF7"
CODE_FILL = "F4F6F8"
BORDER_COLOR = "B7C9D6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER_COLOR) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_doc_core_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.color.rgb = TITLE_COLOR
        style.font.bold = True
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    for list_style_name in ("List Bullet", "List Number"):
        style = doc.styles[list_style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(2)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("RestaurantOS Documentation Pack")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 116, 139)


def add_cover(doc: Document, title: str, subtitle: str, source_name: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("RestaurantOS")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(91, 106, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(35)
    run = p.add_run(subtitle)
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(69, 90, 110)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    rows = [
        ("Document status", "UAT / operational documentation"),
        ("Prepared for", "RestaurantOS stakeholders, administrators, and operations users"),
        ("Source", source_name or "RestaurantOS documentation set"),
        ("Format", "Microsoft Word .docx"),
    ]
    for row, values in zip(table.rows, rows):
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if idx == 0:
                set_cell_shading(cell, LIGHT_FILL)
                run = cell.paragraphs[0].add_run(value)
                run.bold = True
            else:
                cell.paragraphs[0].add_run(value)
    doc.add_page_break()


def add_inline_runs(paragraph, text: str, bold_default: bool = False) -> None:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("`", "")
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_default
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            bold = True
        run = paragraph.add_run(part)
        run.bold = bold


def split_table_row(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip().replace("<br>", "\n") for cell in line.split("|")]


def is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if "|" not in stripped:
        return False
    return all(set(part.strip()) <= {"-", ":"} and "-" in part for part in split_table_row(stripped))


def add_markdown_table(doc: Document, lines: Sequence[str]) -> None:
    header = split_table_row(lines[0])
    body_lines = list(lines[2:]) if len(lines) > 1 and is_table_separator(lines[1]) else list(lines[1:])
    rows = [split_table_row(line) for line in body_lines]
    col_count = max(len(header), *(len(row) for row in rows)) if rows else len(header)
    table = doc.add_table(rows=1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx in range(col_count):
        cell = hdr.cells[idx]
        set_cell_shading(cell, TABLE_HEADER_FILL)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline_runs(p, header[idx] if idx < len(header) else "", bold_default=True)
    for row_values in rows:
        row = table.add_row()
        for idx in range(col_count):
            cell = row.cells[idx]
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, row_values[idx] if idx < len(row_values) else "")
    doc.add_paragraph()


def add_code_block(doc: Document, code: str, language: str = "") -> None:
    label = "Diagram definition" if language.lower() == "mermaid" else "Code block"
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(2)
    run = caption.add_run(label)
    run.bold = True
    run.font.color.rgb = TITLE_COLOR
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(45, 55, 72)
    doc.add_paragraph()


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    set_cell_border(cell, ACCENT_COLOR)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text)
    doc.add_paragraph()


def add_markdown_content(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        fence = re.match(r"^```(\w*)", stripped)
        if fence:
            language = fence.group(1)
            block: List[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, "\n".join(block), language)
            i += 1
            continue

        if "|" in stripped and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            table_lines = [line]
            i += 1
            while i < len(lines) and "|" in lines[i].strip() and lines[i].strip():
                table_lines.append(lines[i].rstrip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            doc.add_heading(heading.group(2), level=level)
            i += 1
            continue

        if stripped.startswith(">"):
            add_callout(doc, stripped.lstrip("> "))
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet.group(1))
            i += 1
            continue

        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, number.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1


def build_document(source_files: Iterable[tuple[str, str, str]], output_path: Path, combined: bool = False) -> None:
    doc = Document()
    set_doc_core_styles(doc)

    if combined:
        add_cover(doc, "RestaurantOS Documentation Pack", "UAT, User Manual, Flow Diagrams, Sequence Diagrams, and System Architecture")
        doc.add_heading("Document Contents", level=1)
        for _, title, subtitle in source_files:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, f"{title} - {subtitle}")
        doc.add_page_break()

    for index, (filename, title, subtitle) in enumerate(source_files):
        if combined:
            if index:
                doc.add_section(WD_SECTION_START.NEW_PAGE)
            doc.add_heading(title, level=1)
            intro = doc.add_paragraph()
            add_inline_runs(intro, subtitle)
        else:
            add_cover(doc, title, subtitle, filename)
        markdown = (DOCS_DIR / filename).read_text(encoding="utf-8")
        add_markdown_content(doc, markdown)

    add_footer(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for filename, title, subtitle in DOC_SOURCES:
        safe_name = Path(filename).stem.lower()
        build_document([(filename, title, subtitle)], OUT_DIR / f"{safe_name}.docx")
    build_document(DOC_SOURCES, OUT_DIR / "RestaurantOS_Documentation_Pack_OnPrem.docx", combined=True)


if __name__ == "__main__":
    main()
